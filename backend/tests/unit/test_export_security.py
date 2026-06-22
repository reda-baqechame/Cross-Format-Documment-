"""Security regressions for exported document bytes."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import UTC, datetime

import fitz
from openpyxl import load_workbook

from docos.model.document import CanonicalDocument, DocumentMeta
from docos.model.nodes import (
    FieldNode,
    RootNode,
    RunNode,
    TableCellNode,
    TableNode,
    TableRowNode,
)
from docos.services.docengine.adapters.pdf import PdfAdapter
from docos.services.docengine.adapters.txt import TxtAdapter
from docos.services.docengine.writers.markup import (
    model_to_csv,
    model_to_html,
    model_to_markdown,
    model_to_rtf,
)
from docos.services.docengine.writers.pdf_writer import write_back_pdf
from docos.services.docengine.writers.xlsx_writer import model_to_xlsx
from docos.services.provenance.validation import validate_export


def _doc(nodes: dict, root_id: str = "root") -> CanonicalDocument:
    now = datetime.now(UTC)
    return CanonicalDocument(
        doc_id="security-doc",
        root_id=root_id,
        nodes=nodes,
        meta=DocumentMeta(
            title="Security export",
            source_format="txt",
            source_mime="text/plain",
            created_at=now,
            modified_at=now,
            page_count=1,
        ),
    )


def test_redacted_field_value_is_removed_and_validation_detects_leak():
    root = RootNode(id="root", children=["field"])
    field = FieldNode(
        id="field",
        parent_id="root",
        field_name="SSN",
        value="123-45-6789",
    )
    doc = _doc({root.id: root, field.id: field})
    doc.redaction.redacted_node_ids.append(field.id)

    html = model_to_html(doc)
    assert b"123-45-6789" not in html
    assert validate_export(doc, "html", html).ok is True

    leaked = b"<html><body>123-45-6789</body></html>"
    report = validate_export(doc, "html", leaked)
    assert report.ok is False
    assert any(f.code == "redaction.recovery" and f.level == "fail" for f in report.findings)


def test_rtf_export_is_valid_and_redaction_safe():
    root = RootNode(id="root", children=["field"])
    field = FieldNode(id="field", parent_id="root", field_name="SSN", value="123-45-6789")
    doc = _doc({root.id: root, field.id: field})
    doc.redaction.redacted_node_ids.append(field.id)

    rtf = model_to_rtf(doc)
    assert rtf[:5] == b"{\\rtf"
    assert rtf[-1:] == b"}"
    assert b"123-45-6789" not in rtf  # redacted content never reaches output


def test_csv_and_xlsx_escape_spreadsheet_formula_payloads():
    formula = '=WEBSERVICE("http://127.0.0.1/leak")'
    root = RootNode(id="root", children=["table"])
    table = TableNode(id="table", parent_id="root", children=["row"], rows=1, cols=1)
    row = TableRowNode(id="row", parent_id="table", children=["cell"], row=0)
    cell = TableCellNode(id="cell", parent_id="row", children=["run"], row=0, col=0)
    run = RunNode(id="run", parent_id="cell", text=formula)
    doc = _doc({n.id: n for n in (root, table, row, cell, run)})

    csv_out = model_to_csv(doc).decode("utf-8")
    cell = next(csv.reader(io.StringIO(csv_out)))[0]
    assert cell == "'" + formula

    xlsx = model_to_xlsx(doc)
    wb = load_workbook(io.BytesIO(xlsx), data_only=False)
    assert wb.worksheets[0]["A1"].value == "'" + formula
    with zipfile.ZipFile(io.BytesIO(xlsx)) as zf:
        package_text = "\n".join(
            zf.read(name).decode("utf-8", "ignore")
            for name in zf.namelist()
            if name.endswith(".xml")
        )
    assert "<f>" not in package_text


def test_html_and_markdown_drop_unsafe_link_schemes():
    doc = TxtAdapter().parse(b"click me")
    run = next(node for node in doc.nodes.values() if node.type == "run")
    run.link_href = "javascript:alert(1)"

    html = model_to_html(doc).decode("utf-8")
    markdown = model_to_markdown(doc).decode("utf-8")
    assert "javascript:" not in html
    assert "javascript:" not in markdown
    assert "<a href" not in html
    assert "[click me]" not in markdown


def test_pdf_writeback_strips_active_links_and_javascript():
    pdf = fitz.open()
    page = pdf.new_page(width=300, height=200)
    page.insert_text((40, 60), "safe pdf text", fontsize=12)
    page.insert_link(
        {
            "kind": fitz.LINK_URI,
            "from": fitz.Rect(20, 20, 140, 80),
            "uri": "javascript:alert(1)",
        }
    )
    data = pdf.tobytes()
    pdf.close()

    doc = PdfAdapter().parse(data)
    out = write_back_pdf(data, doc)
    reopened = fitz.open(stream=out, filetype="pdf")
    try:
        assert all(not page.get_links() for page in reopened)
    finally:
        reopened.close()
    assert b"javascript:alert" not in out.lower()
    assert validate_export(doc, "pdf", out).ok is True
